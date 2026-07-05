# -*- coding: utf-8 -*-
"""生成系统报告 Word 文档"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

IMG = r"C:\Users\Administrator\repos\hr-system\report\images"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

t = doc.add_heading("公司人事管理系统 — 项目报告", 0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("前端：Vue 2.7 + Element UI　　后端：Django 6 + Django REST Framework　　数据库：SQLite")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def img(path, width=6.5, caption=None):
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(os.path.join(IMG, path), width=Inches(width))
    if caption:
        c = doc.add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].font.size = Pt(9)


doc.add_heading("一、系统概述", 1)
doc.add_paragraph(
    "本系统是一套面向中小型公司的人事管理系统，覆盖公司日常人事工作的核心场景："
    "组织架构（部门、职位）管理、员工档案管理、日常考勤记录、请假申请与审批、以及月度薪资核算。"
    "系统采用前后端分离架构：前端使用 Vue 2.7 配合 Element UI 组件库构建单页应用（SPA），"
    "通过 Axios 以 RESTful API 的方式与后端交互；后端基于 Django 6 与 Django REST Framework（DRF）"
    "提供数据接口，使用 Django ORM 访问 SQLite 数据库。"
)

doc.add_heading("二、系统架构图", 1)
doc.add_paragraph(
    "系统分为三层：前端展示层（Vue + Element UI，Vite 开发服务器运行于 5173 端口，"
    "开发环境下通过 Vite 代理将 /api 请求转发至后端 8000 端口）；"
    "后端服务层（Django + DRF，由 DefaultRouter 生成 RESTful 路由，ViewSet 处理 CRUD 与请假审批动作，"
    "Serializer 负责数据序列化与校验）；数据持久层（SQLite，通过 Django ORM 访问）。"
)
img("architecture.png", caption="图 1  系统架构图")

doc.add_heading("三、数据库设计（ER 图）", 1)
doc.add_paragraph(
    "系统共 6 个实体：部门（Department）、职位（Position）、员工（Employee）、"
    "考勤（Attendance）、请假（LeaveRequest）、薪资（Salary）。主要关系如下："
)
for line in [
    "部门 1:N 职位 —— 每个部门下设多个职位；",
    "部门 1:N 员工、职位 1:N 员工 —— 员工归属于某部门并担任某职位（外键保护删除，有员工的部门/职位不可删除）；",
    "员工 1:N 考勤 —— 每名员工每天至多一条考勤记录（唯一约束 employee+date）；",
    "员工 1:N 请假 —— 请假单含类型、起止日期、原因及审批状态（待审批/已批准/已拒绝）；",
    "员工 1:N 薪资 —— 每名员工每月至多一条薪资记录（唯一约束 employee+month），实发工资=基本工资+奖金-扣款。",
]:
    doc.add_paragraph(line, style="List Bullet")
img("er_diagram.png", caption="图 2  ER 图")

doc.add_heading("四、功能说明与运行截图", 1)

doc.add_heading("4.1 首页概览", 2)
doc.add_paragraph(
    "展示在职员工总数、部门数量、待审批请假数三项统计卡片，并以进度条形式展示各部门人数分布，"
    "数据来自后端 /api/dashboard/ 统计接口。"
)
img("01_dashboard.png", caption="图 3  首页概览")

doc.add_heading("4.2 员工管理", 2)
doc.add_paragraph(
    "支持员工的增删改查：可按姓名模糊搜索、按部门筛选；新增/编辑对话框中选择部门后职位下拉框会联动过滤，"
    "只显示该部门下的职位；员工状态区分在职/离职。"
)
img("02_employees.png", caption="图 4  员工列表与查询")
img("03_employee_add.png", caption="图 5  新增员工对话框")
img("04_employee_saved.png", caption="图 6  新增员工保存成功")

doc.add_heading("4.3 部门管理", 2)
doc.add_paragraph(
    "维护公司部门信息（名称、描述），列表实时统计各部门员工人数；删除有在职员工的部门时后端外键保护会拒绝并给出提示。"
)
img("05_departments.png", caption="图 7  部门管理")

doc.add_heading("4.4 职位管理", 2)
doc.add_paragraph("维护各部门下的职位与职级（如 P5、M1），支持新增、编辑、删除。")
img("06_positions.png", caption="图 8  职位管理")

doc.add_heading("4.5 考勤管理", 2)
doc.add_paragraph(
    "记录员工每日签到/签退时间，考勤类型分为正常、迟到、早退、缺勤并以不同颜色标签展示；"
    "同一员工同一天仅允许一条记录。"
)
img("07_attendance.png", caption="图 9  考勤管理")

doc.add_heading("4.6 请假审批", 2)
doc.add_paragraph(
    "员工可发起请假（年假/病假/事假/其他），管理员对待审批记录执行批准或拒绝操作，"
    "对应后端 /api/leaves/{id}/approve/ 与 /reject/ 动作接口。下图为批准操作生效后的界面。"
)
img("08_leaves.png", caption="图 10  请假审批（批准操作生效）")

doc.add_heading("4.7 薪资管理", 2)
doc.add_paragraph(
    "按月录入员工基本工资、奖金与扣款，系统自动计算实发工资（基本工资+奖金-扣款）；"
    "同一员工同一月份仅允许一条记录。"
)
img("09_salaries.png", caption="图 11  薪资管理")

doc.add_heading("五、主要 API 接口", 1)
table = doc.add_table(rows=1, cols=3)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "接口"
hdr[1].text = "方法"
hdr[2].text = "说明"
for row in [
    ("/api/dashboard/", "GET", "首页统计数据"),
    ("/api/departments/", "GET/POST/PUT/DELETE", "部门 CRUD"),
    ("/api/positions/?department=", "GET/POST/PUT/DELETE", "职位 CRUD，可按部门过滤"),
    ("/api/employees/?name=&department=&status=", "GET/POST/PUT/DELETE", "员工 CRUD，支持搜索过滤"),
    ("/api/attendances/", "GET/POST/PUT/DELETE", "考勤 CRUD"),
    ("/api/leaves/", "GET/POST/PUT/DELETE", "请假 CRUD"),
    ("/api/leaves/{id}/approve/ 及 /reject/", "POST", "请假审批：批准 / 拒绝"),
    ("/api/salaries/", "GET/POST/PUT/DELETE", "薪资 CRUD"),
]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text, cells[2].text = row

doc.add_heading("六、运行方式", 1)
doc.add_paragraph("后端（8000 端口）：", style="List Bullet")
doc.add_paragraph("cd backend && pip install django djangorestframework django-cors-headers\n"
                  "python manage.py migrate && python seed.py（初始化示例数据）\n"
                  "python manage.py runserver 0.0.0.0:8000")
doc.add_paragraph("前端（5173 端口）：", style="List Bullet")
doc.add_paragraph("cd frontend && npm install && npm run dev，浏览器访问 http://localhost:5173/")

doc.save(r"C:\Users\Administrator\repos\hr-system\report\公司人事管理系统报告.docx")
print("report done")
